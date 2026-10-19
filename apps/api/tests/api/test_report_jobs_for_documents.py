"""Queued report routes, followed through to the documents they file.

``POST /finance/inventory-valuation/export`` and ``GET /packing/{id}/packing-list``
and the purchase-return debit-note route hand work to Celery and answer with a
job id that ``GET /reports/jobs/{id}`` resolves. This walks the journey a client
actually takes: queue, poll and read QUEUED, let the worker run, poll again and
read COMPLETED with a download link, and find the document row the link points
at. A route that returned a job id nothing could resolve, or a worker that
finished without completing the job row, passes every service-level test and
leaves the client polling for ever.

Committed for real, in a private schema: the worker opens its own session, so
rows left inside the ``session`` fixture's rolled-back transaction would be
invisible to it. The task is run the way Celery's ``AsyncTask`` base runs it -
driving a fresh event loop with ``run_until_complete`` - which cannot happen on
the loop this test is already on, so it runs on a worker thread; the engine is
``NullPool``, so the connections it opens belong to that loop alone. The worker
is pointed at the private schema by replacing its ``session_scope``, and
``send_task`` is stubbed: the broker is not what is under test.
"""

from __future__ import annotations

import asyncio
import uuid
from collections.abc import AsyncIterator, Iterator
from contextlib import asynccontextmanager
from dataclasses import replace
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from typing import Any

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from src.core.config import settings
from src.core.database import get_db, transaction
from src.core.deps import Principal, TenantContext, get_current_principal
from src.core.models import Base
from src.core.permissions import Permission
from src.modules.customers.models import Customer, CustomerType
from src.modules.documents.models import Document, DocumentCategory
from src.modules.documents.service import resolve_storage_path
from src.modules.inventory import service as inventory_service
from src.modules.inventory.models import TransactionType
from src.modules.job_work.models import (
    JobWorkChallan,
    JobWorkChallanLine,
    JobWorkOrder,
    JobWorkOrderLine,
    JobWorkOrderStatus,
    JobWorkProcess,
)
from src.modules.packing import service as packing_service
from src.modules.packing.models import PackingOrderStatus
from src.modules.packing.schemas import PackingOrderCreate, PackingOrderLineIn
from src.modules.products.models import ItemGroup, PriceType
from src.modules.reports import service as reports_service
from src.modules.reports.models import ReportKey
from src.modules.reports.repository import ReportSnapshotRepository
from src.modules.warehouses.models import WarehouseType
from src.workers import celery_app as celery_module
from src.workers.tasks import reports as report_tasks
from tests import seed_support
from tests.schema_support import prepare_metadata, report_database_setup_failure

pytestmark = [pytest.mark.api, pytest.mark.integration]

# ``create_all`` resolves foreign keys against the metadata, so every module has
# to be imported first - see tests/schema_support.
prepare_metadata()


def _schema_engine(schema: str):
    """Engine pinned to one schema, one connection per session (NullPool)."""
    return create_async_engine(
        settings.database_url,
        poolclass=NullPool,
        connect_args={"server_settings": {"search_path": f"{schema},public"}},
    )


def _factory(engine: Any) -> async_sessionmaker[AsyncSession]:
    return async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False, autoflush=False)


@pytest.fixture(scope="module")
def _schema() -> Iterator[tuple[str, uuid.UUID, uuid.UUID]]:
    """A private schema with one company and one user, dropped afterwards."""
    schema = f"it_report_jobs_{uuid.uuid4().hex[:10]}"
    company_id = uuid.uuid4()
    user_id = uuid.uuid4()

    async def setup() -> None:
        admin = create_async_engine(settings.database_url, poolclass=NullPool)
        try:
            async with admin.begin() as conn:
                await conn.execute(text(f'CREATE SCHEMA "{schema}"'))
                await conn.execute(text("CREATE EXTENSION IF NOT EXISTS pg_trgm"))
        finally:
            await admin.dispose()

        engine = _schema_engine(schema)
        try:
            async with engine.begin() as conn:
                await conn.run_sync(
                    lambda sync_conn: Base.metadata.create_all(sync_conn, checkfirst=False)
                )
            async with _factory(engine)() as session, session.begin():
                await seed_support.seed_company(session, company_id=company_id, prefix="RPT")
                await seed_support.seed_user(
                    session,
                    company_id=company_id,
                    user_id=user_id,
                    email=f"reports-{uuid.uuid4().hex[:8]}@atcraft.test",
                )
        finally:
            await engine.dispose()

    async def teardown() -> None:
        admin = create_async_engine(settings.database_url, poolclass=NullPool)
        try:
            async with admin.begin() as conn:
                await conn.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
        finally:
            await admin.dispose()

    try:
        asyncio.run(setup())
    except Exception as exc:  # pragma: no cover - environment, not a test failure
        report_database_setup_failure(exc)

    yield schema, company_id, user_id
    asyncio.run(teardown())


@pytest_asyncio.fixture
async def engine(_schema: tuple[str, uuid.UUID, uuid.UUID]) -> AsyncIterator[Any]:
    schema, _, _ = _schema
    eng = _schema_engine(schema)
    try:
        yield eng
    finally:
        await eng.dispose()


@pytest.fixture
def tenant(_schema: tuple[str, uuid.UUID, uuid.UUID]) -> TenantContext:
    _, company_id, user_id = _schema
    return TenantContext(company_id=company_id, user_id=user_id)


@pytest.fixture
def principal(_schema: tuple[str, uuid.UUID, uuid.UUID]) -> Principal:
    _, company_id, user_id = _schema
    return Principal(
        user_id=user_id,
        company_id=company_id,
        email="reports@atcraft.test",
        full_name="Reports Tester",
        is_superuser=False,
        permissions=frozenset(str(p) for p in Permission),
        role_codes=("TEST_ADMIN",),
    )


@pytest_asyncio.fixture
async def client(
    engine: Any, principal: Principal, monkeypatch: pytest.MonkeyPatch, tmp_path
) -> AsyncIterator[AsyncClient]:
    """The app over ASGI, its request sessions and the worker's sessions both
    bound to the private schema, the broker stubbed and files kept in tmp_path."""
    from src.main import app

    factory = _factory(engine)

    async def request_session() -> AsyncIterator[AsyncSession]:
        async with factory() as session:
            yield session

    @asynccontextmanager
    async def worker_session() -> AsyncIterator[AsyncSession]:
        async with factory() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    def accepted_by_the_broker(*_args: Any, **_kwargs: Any) -> SimpleNamespace:
        return SimpleNamespace(id=f"stub-{uuid.uuid4().hex}")

    monkeypatch.setattr(settings, "storage_path", tmp_path)
    monkeypatch.setattr(celery_module.celery_app, "send_task", accepted_by_the_broker)
    monkeypatch.setattr(report_tasks, "session_scope", worker_session)
    app.dependency_overrides[get_db] = request_session
    app.dependency_overrides[get_current_principal] = lambda: principal
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as http:
            yield http
    finally:
        app.dependency_overrides.clear()


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------


def _code(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8].upper()}"


def _prefix() -> str:
    return settings.api_v1_prefix.rstrip("/")


async def _run_worker(task: Any, **kwargs: str) -> dict[str, Any]:
    """Run a Celery task inline, off the test's own event loop."""
    return await asyncio.to_thread(task, **kwargs)


async def _job(client: AsyncClient, status_url: str) -> dict[str, Any]:
    response = await client.get(status_url)
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def _filed_document(engine: Any, document_id: str) -> Document:
    async with _factory(engine)() as reader:
        document = await reader.get(Document, uuid.UUID(document_id))
    assert document is not None, "the job reports a document that was never filed"
    return document


def _assert_queued(job: dict[str, Any], *, job_id: str, report_key: str) -> None:
    assert job["id"] == job_id
    assert job["status"] == "QUEUED"
    assert job["report_key"] == report_key
    assert job["document_id"] is None
    assert job["download_url"] is None


def _assert_completed(job: dict[str, Any], *, job_id: str, report_key: str) -> str:
    assert job["id"] == job_id
    assert job["status"] == "COMPLETED", job
    assert job["report_key"] == report_key
    assert job["error_message"] is None
    document_id = job["document_id"]
    assert document_id is not None
    assert job["download_url"] == (
        reports_service.job_download_url(job_id)
        if report_key in {"INVENTORY_VALUATION", "JOB_WORK_CHALLAN", "PURCHASE_RETURN_DEBIT_NOTE"}
        else reports_service.document_download_url(document_id)
    )
    return str(document_id)


async def _stock_a_position(engine: Any, tenant: TenantContext) -> None:
    """One valued position, so the workbook has a row to carry."""
    async with _factory(engine)() as session, transaction(session):
        warehouse = await seed_support.seed_warehouse(
            session, company_id=tenant.company_id, code=_code("WH")
        )
        item = await seed_support.seed_item(
            session,
            company_id=tenant.company_id,
            code=_code("RM"),
            gst_rate=Decimal("18"),
            standard_cost=Decimal("10"),
        )
        await inventory_service.post_movement(
            session,
            tenant,
            item_id=item.id,
            warehouse_id=warehouse.id,
            quantity=Decimal("25"),
            transaction_type=TransactionType.OPENING_BALANCE,
            reference_type="opening",
            reference_id=None,
            reference_code="OPENING",
            reason="report job fixture opening stock",
            unit_cost=Decimal("10"),
        )


async def _packing_in_progress(engine: Any, tenant: TenantContext) -> uuid.UUID:
    """A packing order on the bench - the earliest state a packing list means anything."""
    async with _factory(engine)() as session:
        async with transaction(session):
            uom = await seed_support.seed_uom(session, company_id=tenant.company_id)
            warehouse = await seed_support.seed_warehouse(
                session,
                company_id=tenant.company_id,
                code=_code("FG"),
                warehouse_type=WarehouseType.FINISHED_GOODS,
            )
            goods = await seed_support.seed_item(
                session,
                company_id=tenant.company_id,
                code=_code("LOCK"),
                base_uom_id=uom.id,
                item_group=ItemGroup.FINISHED_GOOD,
                gst_rate=Decimal("18"),
                standard_cost=Decimal("60"),
                sale_price=Decimal("100"),
                can_sell=True,
            )
            customer = Customer(
                company_id=tenant.company_id,
                code=_code("CU"),
                name="Hardware Distributors",
                customer_type=CustomerType.CUSTOMER,
                price_list_type=PriceType.SALE,
                credit_days=30,
                is_active=True,
                is_blocked=False,
            )
            session.add(customer)
            await session.flush()

        order, _ = await packing_service.create_order(
            session,
            tenant,
            PackingOrderCreate(
                customer_id=customer.id,
                warehouse_id=warehouse.id,
                planned_date=date(2026, 9, 1),
                lines=[
                    PackingOrderLineIn(
                        item_id=goods.id, quantity_to_pack=Decimal("10"), uom_id=uom.id
                    )
                ],
            ),
        )
        started, _ = await packing_service.start_packing(
            session, tenant, order.id, reason="Bench is free"
        )
        assert PackingOrderStatus(started.status) is PackingOrderStatus.IN_PROGRESS
        return order.id


# -----------------------------------------------------------------------------
# The journeys
# -----------------------------------------------------------------------------


async def _issued_job_work_challan(engine: Any, tenant: TenantContext) -> uuid.UUID:
    """An issued consignment with real parties and goods for the PDF worker."""
    async with _factory(engine)() as session, transaction(session):
        warehouse = await seed_support.seed_warehouse(
            session, company_id=tenant.company_id, code=_code("WH")
        )
        vendor = await seed_support.seed_vendor(
            session, company_id=tenant.company_id, code=_code("VEN")
        )
        item = await seed_support.seed_item(
            session, company_id=tenant.company_id, code=_code("RAW")
        )
        order = JobWorkOrder(
            company_id=tenant.company_id,
            code=_code("JW"),
            vendor_id=vendor.id,
            process_type=JobWorkProcess.PVD,
            expected_return_date=date(2026, 10, 30),
            warehouse_id=warehouse.id,
            return_warehouse_id=warehouse.id,
            status=JobWorkOrderStatus.ISSUED,
        )
        session.add(order)
        await session.flush()
        line = JobWorkOrderLine(
            company_id=tenant.company_id,
            job_work_order_id=order.id,
            line_number=1,
            item_id=item.id,
            uom_id=item.base_uom_id,
            quantity_to_issue=Decimal("10"),
            issued_quantity=Decimal("10"),
            expected_output_item_id=item.id,
            expected_output_quantity=Decimal("10"),
        )
        session.add(line)
        await session.flush()
        challan = JobWorkChallan(
            company_id=tenant.company_id,
            code=_code("JWC"),
            job_work_order_id=order.id,
            vendor_id=vendor.id,
            challan_date=date(2026, 10, 16),
            expected_return_date=date(2026, 10, 30),
            total_quantity=Decimal("10"),
            total_value=Decimal("1000"),
        )
        session.add(challan)
        await session.flush()
        session.add(
            JobWorkChallanLine(
                company_id=tenant.company_id,
                challan_id=challan.id,
                job_work_order_line_id=line.id,
                line_number=1,
                item_id=item.id,
                uom_id=item.base_uom_id,
                quantity=Decimal("10"),
                unit_value=Decimal("100"),
                total_value=Decimal("1000"),
            )
        )
        return challan.id


async def test_job_work_view_holder_can_retrieve_only_its_own_challan_export(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.core.clock import utcnow
    from src.main import app

    operator = replace(principal, permissions=frozenset({str(Permission.JOB_WORK_VIEW)}))
    app.dependency_overrides[get_current_principal] = lambda: operator
    challan_id = await _issued_job_work_challan(engine, tenant)
    accepted = await client.post(f"{_prefix()}/job-work/challans/{challan_id}/pdf")
    assert accepted.status_code == 202, accepted.text
    job_id = accepted.json()["data"]["job_id"]
    status_url = accepted.json()["data"]["status_url"]
    download_url = reports_service.job_download_url(job_id)
    queued = await _job(client, status_url)
    _assert_queued(queued, job_id=job_id, report_key="JOB_WORK_CHALLAN")
    assert (await client.get(download_url)).status_code == 404

    outcome = await _run_worker(
        report_tasks.generate_job_work_challan,
        job_work_challan_id=str(challan_id),
        snapshot_id=job_id,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
    )
    assert outcome["status"] == "COMPLETED", outcome
    completed = await _job(client, status_url)
    document_id = _assert_completed(completed, job_id=job_id, report_key="JOB_WORK_CHALLAN")
    async with _factory(engine)() as session:
        snapshot = await ReportSnapshotRepository(session, tenant).get_or_404(uuid.UUID(job_id))
    downloaded = await client.get(completed["download_url"])
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.content.startswith(b"%PDF-")
    assert downloaded.headers["x-content-type-options"] == "nosniff"

    # Originating-job access cannot be used for general reports or arbitrary
    # document ids, even by the same operator who generated this PDF.
    assert (await client.get(f"{_prefix()}/reports/catalogue")).status_code == 403
    assert (await client.get(f"{_prefix()}/documents")).status_code == 403
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 403
    substituted = await client.get(download_url, params={"document_id": str(uuid.uuid4())})
    assert substituted.status_code == 200
    assert substituted.content == downloaded.content

    for intruder in (
        replace(operator, user_id=uuid.uuid4()),
        replace(operator, company_id=uuid.uuid4()),
        replace(operator, permissions=frozenset({str(Permission.INVENTORY_VALUATION_VIEW)})),
    ):
        app.dependency_overrides[get_current_principal] = lambda user=intruder: user
        assert (await client.get(status_url)).status_code == 404
        assert (await client.get(download_url)).status_code == 404

    app.dependency_overrides[get_current_principal] = lambda: replace(
        operator, permissions=frozenset()
    )
    assert (await client.get(status_url)).status_code == 403
    assert (await client.get(download_url)).status_code == 403

    for outsider in (
        replace(principal, vendor_id=uuid.uuid4()),
        replace(principal, customer_id=uuid.uuid4()),
    ):
        assert not reports_service._owns_job_work_challan_export(snapshot, outsider)
        app.dependency_overrides[get_current_principal] = lambda user=outsider: user
        assert (await client.get(status_url)).status_code == 403
        assert (await client.get(download_url)).status_code == 403

    # The exception is for the exact document-only challan key, even when a
    # different job is owned by this operator and names the very same document.
    async with _factory(engine)() as session, transaction(session):
        unrelated = reports_service.new_job_snapshot(ReportKey.PACKING_LIST, {}, now=utcnow())
        unrelated.status = reports_service.ReportJobStatus.COMPLETED
        unrelated.document_id = uuid.UUID(document_id)
        unrelated.generated_at = utcnow()
        ReportSnapshotRepository(session, tenant).add(unrelated)
        await session.flush()
        unrelated_id = unrelated.id
    app.dependency_overrides[get_current_principal] = lambda: operator
    assert (await client.get(reports_service.job_status_url(unrelated_id))).status_code == 404
    assert (await client.get(reports_service.job_download_url(unrelated_id))).status_code == 404

    # Existing internal report readers can still share a colleague's job, but
    # their download continues to require the general document-view right.
    reader = replace(
        principal,
        user_id=uuid.uuid4(),
        permissions=frozenset({str(Permission.REPORT_VIEW), str(Permission.DOCUMENT_VIEW)}),
    )
    app.dependency_overrides[get_current_principal] = lambda: reader
    assert (await client.get(status_url)).status_code == 200
    assert (await client.get(download_url)).status_code == 200


async def test_purchase_return_view_holder_can_read_only_its_own_debit_note_job(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.core.clock import utcnow
    from src.main import app

    operator = replace(principal, permissions=frozenset({str(Permission.PURCHASE_RETURN_VIEW)}))
    async with _factory(engine)() as session, transaction(session):
        snapshot = reports_service.new_job_snapshot(
            ReportKey.PURCHASE_RETURN_DEBIT_NOTE,
            {"purchase_return_id": str(uuid.uuid4())},
            now=utcnow(),
        )
        ReportSnapshotRepository(session, tenant).add(snapshot)
        await session.flush()
        job_id = snapshot.id

    app.dependency_overrides[get_current_principal] = lambda: operator
    status_url = reports_service.job_status_url(job_id)
    response = await client.get(status_url)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["report_key"] == "PURCHASE_RETURN_DEBIT_NOTE"

    intruder = replace(operator, user_id=uuid.uuid4())
    app.dependency_overrides[get_current_principal] = lambda: intruder
    assert (await client.get(status_url)).status_code == 404

    app.dependency_overrides[get_current_principal] = lambda: replace(
        operator, permissions=frozenset()
    )
    assert (await client.get(status_url)).status_code == 403


async def test_purchase_return_debit_note_worker_files_audited_pdf_and_scoped_download(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.core.clock import utcnow
    from src.main import app
    from src.modules.audit.models import AuditAction, AuditLog
    from src.modules.procurement import service as procurement_service
    from src.modules.procurement.models import PurchaseOrderStatus
    from src.modules.procurement.schemas import (
        GoodsReceiptCreate,
        GoodsReceiptLineIn,
        PurchaseOrderCreate,
        PurchaseOrderLineIn,
        PurchaseReturnCreate,
        PurchaseReturnLineIn,
    )

    async with _factory(engine)() as session:
        uom = await seed_support.seed_uom(session, company_id=tenant.company_id)
        warehouse = await seed_support.seed_warehouse(
            session, company_id=tenant.company_id, code=_code("WH")
        )
        item = await seed_support.seed_item(
            session,
            company_id=tenant.company_id,
            code=_code("RM"),
            base_uom_id=uom.id,
            purchase_uom_id=uom.id,
            gst_rate=Decimal("18"),
            standard_cost=Decimal("50"),
            can_purchase=True,
        )
        vendor = await seed_support.seed_vendor(
            session,
            company_id=tenant.company_id,
            code=_code("VEN"),
            is_approved=True,
            approved_by_id=tenant.user_id,
            approved_at=utcnow(),
            currency="INR",
            state_code="27",
        )
        order, order_lines, _ = await procurement_service.create_purchase_order(
            session,
            tenant,
            principal,
            PurchaseOrderCreate(
                vendor_id=vendor.id,
                warehouse_id=warehouse.id,
                lines=[
                    PurchaseOrderLineIn(
                        line_number=1,
                        item_id=item.id,
                        uom_id=uom.id,
                        quantity=Decimal("10"),
                        unit_price=Decimal("100"),
                    )
                ],
            ),
        )
        await procurement_service.transition_purchase_order(
            session,
            tenant,
            principal,
            order.id,
            PurchaseOrderStatus.PENDING_APPROVAL,
            reason="submit debit-note PDF fixture",
        )
        await procurement_service.transition_purchase_order(
            session,
            tenant,
            replace(principal, is_superuser=True),
            order.id,
            PurchaseOrderStatus.APPROVED,
            reason="approve debit-note PDF fixture",
        )
        await procurement_service.transition_purchase_order(
            session,
            tenant,
            principal,
            order.id,
            PurchaseOrderStatus.SENT,
            reason="send debit-note PDF fixture",
        )
        receipt, _ = await procurement_service.create_goods_receipt(
            session,
            tenant,
            GoodsReceiptCreate(
                purchase_order_id=order.id,
                lines=[
                    GoodsReceiptLineIn(
                        line_number=1,
                        purchase_order_line_id=order_lines[0].id,
                        received_quantity=Decimal("10"),
                        accepted_quantity=Decimal("10"),
                    )
                ],
            ),
        )
        posted_receipt = await procurement_service.post_goods_receipt(
            session, tenant, principal, receipt.id
        )
        purchase_return, _ = await procurement_service.create_purchase_return(
            session,
            tenant,
            PurchaseReturnCreate(
                goods_receipt_id=posted_receipt.receipt.id,
                reason="Debit-note worker test return",
                debit_note_number="DN/WORKER/1",
                lines=[
                    PurchaseReturnLineIn(
                        line_number=1,
                        goods_receipt_line_id=posted_receipt.receipt.lines[0].id,
                        quantity=Decimal("2"),
                    )
                ],
            ),
        )
        await procurement_service.post_purchase_return(
            session, tenant, principal, purchase_return.id
        )
        purchase_return_id = purchase_return.id

    accepted = await client.post(
        f"{_prefix()}/procurement/purchase-returns/{purchase_return_id}/debit-note-pdf"
    )
    assert accepted.status_code == 202, accepted.text
    handed_off = accepted.json()["data"]
    job_id = handed_off["job_id"]
    _assert_queued(
        await _job(client, handed_off["status_url"]),
        job_id=job_id,
        report_key="PURCHASE_RETURN_DEBIT_NOTE",
    )

    outcome = await _run_worker(
        report_tasks.generate_purchase_return_debit_note,
        purchase_return_id=str(purchase_return_id),
        snapshot_id=job_id,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
    )
    assert outcome["status"] == "COMPLETED", outcome
    status = await _job(client, handed_off["status_url"])
    document_id = _assert_completed(status, job_id=job_id, report_key="PURCHASE_RETURN_DEBIT_NOTE")
    document = await _filed_document(engine, document_id)
    assert document.company_id == tenant.company_id
    assert document.entity_type == report_tasks.PURCHASE_RETURN_ENTITY
    assert document.entity_id == purchase_return_id
    assert document.is_public is False
    assert document.expires_at is None
    assert document.size_bytes > 0
    assert resolve_storage_path(document.storage_key).read_bytes().startswith(b"%PDF")

    restricted_operator = replace(
        principal, permissions=frozenset({str(Permission.PURCHASE_RETURN_VIEW)})
    )
    app.dependency_overrides[get_current_principal] = lambda: restricted_operator
    download = await client.get(status["download_url"])
    assert download.status_code == 200, download.text
    assert download.content.startswith(b"%PDF")

    async with _factory(engine)() as session:
        audit = await session.scalar(
            select(AuditLog).where(
                AuditLog.company_id == tenant.company_id,
                AuditLog.entity_type == report_tasks.PURCHASE_RETURN_ENTITY,
                AuditLog.entity_id == purchase_return_id,
                AuditLog.action == AuditAction.DATA_EXPORT,
            )
        )
    assert audit is not None
    assert audit.created_by_id == tenant.user_id


async def test_the_inventory_valuation_export_is_a_job_that_files_a_workbook(
    client: AsyncClient, engine: Any, tenant: TenantContext
) -> None:
    await _stock_a_position(engine, tenant)

    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202, accepted.text
    handed_off = accepted.json()["data"]
    job_id = handed_off["job_id"]
    assert handed_off["status_url"] == reports_service.job_status_url(job_id)

    _assert_queued(
        await _job(client, handed_off["status_url"]),
        job_id=job_id,
        report_key="INVENTORY_VALUATION",
    )

    outcome = await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=job_id,
    )
    assert outcome["status"] == "COMPLETED", outcome

    document_id = _assert_completed(
        await _job(client, handed_off["status_url"]),
        job_id=job_id,
        report_key="INVENTORY_VALUATION",
    )
    document = await _filed_document(engine, document_id)
    assert document.company_id == tenant.company_id
    assert document.category == DocumentCategory.REPORT
    assert document.file_name.endswith(".xlsx")
    assert document.size_bytes > 0
    assert resolve_storage_path(document.storage_key).is_file()


async def test_restricted_valuation_exporter_can_watch_and_download_only_its_own_export(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.main import app

    exporter = replace(
        principal,
        permissions=frozenset(
            {str(Permission.INVENTORY_VALUATION_VIEW), str(Permission.REPORT_EXPORT)}
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: exporter
    await _stock_a_position(engine, tenant)
    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202, accepted.text
    job_id = accepted.json()["data"]["job_id"]
    status_url = accepted.json()["data"]["status_url"]
    download_url = reports_service.job_download_url(job_id)
    assert (await client.get(status_url)).status_code == 200
    assert (await client.get(download_url)).status_code == 404  # no file until completion

    outcome = await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=job_id,
    )
    assert outcome["status"] == "COMPLETED", outcome
    job = await _job(client, status_url)
    assert job["download_url"] == download_url
    downloaded = await client.get(download_url)
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.content.startswith(b"PK")  # actual generated workbook
    assert downloaded.headers["x-content-type-options"] == "nosniff"
    assert int(downloaded.headers["content-length"]) == len(downloaded.content)
    assert "attachment" in downloaded.headers["content-disposition"]

    # The exceptional right is bound to this server-issued job, not broad file
    # or catalogue access; arbitrary document ids cannot replace that binding.
    assert (await client.get(f"{_prefix()}/reports/catalogue")).status_code == 403
    assert (
        await client.get(f"{_prefix()}/documents/{job['document_id']}/download")
    ).status_code == 403
    substituted = await client.get(download_url, params={"document_id": str(uuid.uuid4())})
    assert substituted.status_code == 200
    assert substituted.content == downloaded.content

    for intruder in (
        replace(exporter, user_id=uuid.uuid4()),
        replace(exporter, company_id=uuid.uuid4()),
        replace(exporter, permissions=frozenset({str(Permission.INVENTORY_VALUATION_VIEW)})),
    ):
        app.dependency_overrides[get_current_principal] = lambda user=intruder: user
        assert (await client.get(status_url)).status_code == 404
        assert (await client.get(download_url)).status_code == 404

    # Even broad grants do not make a vendor an internal report principal.
    vendor = replace(principal, vendor_id=uuid.uuid4())
    app.dependency_overrides[get_current_principal] = lambda: vendor
    assert (await client.get(status_url)).status_code == 403
    assert (await client.get(download_url)).status_code == 403

    # Sharing a valuation export requires its specific right as well as the
    # general report and document rights.
    reader = replace(
        principal,
        user_id=uuid.uuid4(),
        permissions=frozenset(
            {
                str(Permission.REPORT_VIEW),
                str(Permission.DOCUMENT_VIEW),
                str(Permission.INVENTORY_VALUATION_VIEW),
            }
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: reader
    assert (await client.get(status_url)).status_code == 200
    assert (await client.get(download_url)).status_code == 200
    app.dependency_overrides[get_current_principal] = lambda: replace(
        reader, permissions=frozenset({str(Permission.REPORT_VIEW)})
    )
    assert (await client.get(status_url)).status_code == 404
    assert (await client.get(download_url)).status_code == 404


async def test_dealer_with_broad_grants_cannot_poll_download_or_queue_finance_export(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from sqlalchemy import func, select

    from src.core.clock import utcnow
    from src.core.errors import NotFoundError
    from src.main import app
    from src.modules.finance.service import VALUATION_TASK_NAME
    from src.modules.reports.models import ReportKey, ReportSnapshot
    from src.modules.reports.repository import ReportSnapshotRepository

    # The dealer is deliberately given internal permissions and owns the job.
    # Auth is overridden, so the explicit route/service portal gates must work.
    dealer = replace(principal, customer_id=uuid.uuid4())
    async with _factory(engine)() as session, session.begin():
        snapshot = reports_service.new_job_snapshot(
            ReportKey.INVENTORY_VALUATION, {"task": VALUATION_TASK_NAME}, now=utcnow()
        )
        ReportSnapshotRepository(session, tenant).add(snapshot)
        await session.flush()
        job_id = snapshot.id
        assert not reports_service._owns_valuation_export(snapshot, dealer)
        with pytest.raises(NotFoundError):
            await reports_service.get_accessible_job(session, tenant, dealer, job_id)

    app.dependency_overrides[get_current_principal] = lambda: dealer
    for url in (
        reports_service.job_status_url(job_id),
        reports_service.job_download_url(job_id),
    ):
        result = await client.get(url)
        assert result.status_code == 403
        assert result.json()["error"]["code"] == "SCOPE_DENIED"

    async with _factory(engine)() as session:
        before = await session.scalar(select(func.count()).select_from(ReportSnapshot))
    queued = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert queued.status_code == 403
    assert queued.json()["error"]["code"] == "SCOPE_DENIED"
    async with _factory(engine)() as session:
        assert await session.scalar(select(func.count()).select_from(ReportSnapshot)) == before


async def test_valuation_permission_does_not_open_catalogue_jobs_or_other_document_jobs(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.core.clock import utcnow
    from src.main import app

    ids = []
    async with _factory(engine)() as session, transaction(session):
        repo = ReportSnapshotRepository(session, tenant)
        for key, params in (
            (ReportKey.INVENTORY_VALUATION, {"filters": {}}),  # same key, different origin
            (ReportKey.PACKING_LIST, {"packing_order_id": str(uuid.uuid4())}),
            (ReportKey.STOCK_AVAILABILITY, {"task": "reports.generate_inventory_valuation"}),
        ):
            snapshot = reports_service.new_job_snapshot(key, params, now=utcnow())
            repo.add(snapshot)
            await session.flush()
            ids.append(snapshot.id)
    exporter = replace(
        principal,
        permissions=frozenset(
            {str(Permission.INVENTORY_VALUATION_VIEW), str(Permission.REPORT_EXPORT)}
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: exporter
    for job_id in ids:
        assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
        assert (await client.get(reports_service.job_download_url(job_id))).status_code == 404

    # Generic access remains available for non-sensitive catalogue reports,
    # but not for valuation jobs.
    app.dependency_overrides[get_current_principal] = lambda: replace(
        principal, permissions=frozenset({str(Permission.REPORT_VIEW)})
    )
    assert (await client.get(reports_service.job_status_url(ids[0]))).status_code == 404
    for job_id in ids[1:]:
        assert (await client.get(reports_service.job_status_url(job_id))).status_code == 200


async def test_restricted_catalogue_report_job_and_document_keep_captured_scope(
    client: AsyncClient, engine: Any, tenant: TenantContext, principal: Principal
) -> None:
    from src.core.clock import utcnow
    from src.main import app

    (allowed, _), (outside, _) = await _stock_two_warehouses(engine, tenant)
    scope_key = reports_service.REPORT_WAREHOUSE_SCOPE_PARAMETER
    parameters = {
        "report_key": str(ReportKey.INVENTORY_VALUATION),
        "filters": {},
        scope_key: [str(allowed)],
    }
    async with _factory(engine)() as session:
        snapshot = reports_service.new_job_snapshot(
            ReportKey.INVENTORY_VALUATION, parameters, now=utcnow()
        )
        ReportSnapshotRepository(session, tenant).add(snapshot)
        await session.flush()
        document = Document(
            company_id=tenant.company_id,
            file_name="catalogue-valuation.xlsx",
            storage_key=f"reports/{tenant.company_id}/catalogue-scope.xlsx",
            mime_type="application/octet-stream",
            size_bytes=1,
            checksum_sha256="a" * 64,
            category=DocumentCategory.REPORT,
            entity_type="report_snapshot",
            entity_id=snapshot.id,
            uploaded_by_id=tenant.user_id,
        )
        session.add(document)
        await session.flush()
        await reports_service.start_snapshot(session, tenant, snapshot.id)
        await reports_service.complete_snapshot(
            session, tenant, snapshot_id=snapshot.id, row_count=1, document_id=document.id
        )
        job_id = snapshot.id
        document_id = document.id
        await session.commit()

    restricted_reader = replace(
        principal,
        allowed_warehouse_ids=(allowed,),
        permissions=frozenset(
            {
                str(Permission.REPORT_VIEW),
                str(Permission.DOCUMENT_VIEW),
                str(Permission.INVENTORY_VALUATION_VIEW),
            }
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: restricted_reader
    assert (await client.get(reports_service.job_status_url(job_id))).status_code == 200
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 200

    outside_reader = replace(restricted_reader, allowed_warehouse_ids=(outside,))
    app.dependency_overrides[get_current_principal] = lambda: outside_reader
    assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404
    listed = await client.get(
        f"{_prefix()}/documents",
        params={"entity_type": "report_snapshot", "entity_id": str(job_id)},
    )
    assert listed.status_code == 200
    assert listed.json()["data"] == []


async def _stock_two_warehouses(engine, tenant):
    """Different values make an overbroad export distinguishable from its page."""
    warehouses = []
    async with _factory(engine)() as session, transaction(session):
        item = await seed_support.seed_item(
            session, company_id=tenant.company_id, code=_code("SCOPE")
        )
        for quantity, cost in [("5", "2"), ("7", "11")]:
            warehouse = await seed_support.seed_warehouse(
                session, company_id=tenant.company_id, code=_code("SCOPE-WH")
            )
            warehouses.append((warehouse.id, warehouse.code))
            await inventory_service.post_movement(
                session,
                tenant,
                item_id=item.id,
                warehouse_id=warehouse.id,
                quantity=Decimal(quantity),
                unit_cost=Decimal(cost),
                transaction_type=TransactionType.OPENING_BALANCE,
                reference_type="opening",
                reference_id=None,
                reference_code="SCOPE-OPEN",
                reason="Valuation scope regression",
            )
    return warehouses


async def test_catalogue_report_worker_uses_and_preserves_captured_warehouse_scope(
    client: AsyncClient,
    engine: Any,
    tenant: TenantContext,
    principal: Principal,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from src.main import app
    from src.modules.reports.models import ReportSnapshot

    (allowed, _), (outside, _) = await _stock_two_warehouses(engine, tenant)
    async with _factory(engine)() as session:
        await reports_service.refresh_inventory_summary(session, tenant)
        await session.commit()
    restricted = replace(principal, allowed_warehouse_ids=(allowed,))
    app.dependency_overrides[get_current_principal] = lambda: restricted
    monkeypatch.setattr(reports_service, "REPORT_INLINE_ROW_LIMIT", 0)
    accepted = await client.get(f"{_prefix()}/reports/inventory-valuation")
    assert accepted.status_code == 202, accepted.text
    job_id = accepted.json()["data"]["job_id"]
    async with _factory(engine)() as session:
        snapshot = await session.get(ReportSnapshot, uuid.UUID(job_id))
        assert snapshot.parameters[reports_service.REPORT_WAREHOUSE_SCOPE_PARAMETER] == [
            str(allowed)
        ]

    outcome = await _run_worker(
        report_tasks.generate_report,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=job_id,
    )
    assert outcome["status"] == "COMPLETED"
    job = await _job(client, reports_service.job_status_url(job_id))
    assert job["row_count"] == 1
    document_id = job["document_id"]
    downloaded = await client.get(reports_service.job_download_url(job_id))
    assert downloaded.status_code == 200
    workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
    assert workbook.active.max_row == 2
    workbook.close()

    outside_reader = replace(restricted, allowed_warehouse_ids=(outside,))
    app.dependency_overrides[get_current_principal] = lambda: outside_reader
    assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
    assert (await client.get(reports_service.job_download_url(job_id))).status_code == 404
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404


@pytest.mark.parametrize("legacy", [False, True])
async def test_valuation_workbook_matches_inline_warehouse_scope_and_rejects_narrowing(
    client, engine, tenant, principal, legacy
):
    from io import BytesIO

    from openpyxl import load_workbook

    from src.main import app
    from src.modules.reports.models import ReportSnapshot
    from src.modules.users.models import User

    (allowed, code), (outside, _) = await _stock_two_warehouses(engine, tenant)
    exporter = replace(
        principal,
        allowed_warehouse_ids=(allowed,),
        permissions=frozenset(
            {str(Permission.INVENTORY_VALUATION_VIEW), str(Permission.REPORT_EXPORT)}
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: exporter
    inline = await client.get(f"{_prefix()}/finance/inventory-valuation")
    assert inline.status_code == 200
    assert inline.json()["data"]["summary"]["row_count"] == 1
    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202
    job_id = accepted.json()["data"]["job_id"]
    scope_key = reports_service.VALUATION_WAREHOUSE_SCOPE_PARAMETER
    async with _factory(engine)() as session, session.begin():
        snapshot = await session.get(ReportSnapshot, uuid.UUID(job_id))
        assert snapshot.parameters[scope_key] == [str(allowed)]
        if legacy:
            snapshot.parameters = {
                key: value for key, value in snapshot.parameters.items() if key != scope_key
            }
            creator = await session.get(User, tenant.user_id)
            creator.allowed_warehouse_ids = [allowed]
    if legacy:
        assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
    outcome = await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=job_id,
        # Arbitrary broker kwargs cannot replace persisted scope OR filters.
        warehouse_id=str(outside),
        allowed_warehouse_ids=[str(outside)],
    )
    assert outcome["status"] == "COMPLETED"
    job = await _job(client, reports_service.job_status_url(job_id))
    document_id = job["document_id"]
    downloaded = await client.get(reports_service.job_download_url(job_id))
    assert downloaded.status_code == 200
    workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
    rows = list(workbook["Valuation"].values)[1:]
    assert len(rows) == 1
    assert rows[0][2] == code
    summary = dict(workbook["Summary"].values)
    assert Decimal(str(summary["Total value"])) == Decimal(
        inline.json()["data"]["summary"]["total_value"]
    )
    workbook.close()
    async with _factory(engine)() as session:
        snapshot = await session.get(ReportSnapshot, uuid.UUID(job_id))
        assert snapshot.parameters[scope_key] == [str(allowed)]
        document = await session.get(Document, uuid.UUID(document_id))
        assert document.entity_type == report_tasks.VALUATION_ENTITY
        assert document.entity_id == snapshot.id

    app.dependency_overrides[get_current_principal] = lambda: replace(
        exporter, allowed_warehouse_ids=(allowed, outside)
    )
    assert (await client.get(reports_service.job_download_url(job_id))).status_code == 200
    for permissions in (exporter.permissions, principal.permissions):
        app.dependency_overrides[get_current_principal] = lambda grants=permissions: replace(
            exporter, allowed_warehouse_ids=(outside,), permissions=grants
        )
        assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
        assert (await client.get(reports_service.job_download_url(job_id))).status_code == 404
        if str(Permission.DOCUMENT_VIEW) in permissions:
            assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404
            assert (
                await client.get(f"{_prefix()}/documents/{document_id}/download")
            ).status_code == 404
            listed = await client.get(f"{_prefix()}/documents")
            assert listed.status_code == 200
            assert document_id not in {row["id"] for row in listed.json()["data"]}
    # Existing unrestricted report/document readers still see their company artifacts.
    app.dependency_overrides[get_current_principal] = lambda: principal
    assert (await client.get(reports_service.job_download_url(job_id))).status_code == 200
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 200
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 200


async def test_valuation_scope_evidence_survives_cleanup_until_document_deletion(
    client, engine, tenant, principal
):
    from datetime import timedelta

    from src.core.clock import utcnow
    from src.main import app
    from src.modules.documents.service import DocumentService
    from src.modules.reports.models import ReportSnapshot

    (allowed, _), (outside, _) = await _stock_two_warehouses(engine, tenant)
    app.dependency_overrides[get_current_principal] = lambda: replace(
        principal, allowed_warehouse_ids=(allowed,)
    )
    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202
    job_id = uuid.UUID(accepted.json()["data"]["job_id"])
    await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=str(job_id),
    )
    async with _factory(engine)() as session, session.begin():
        snapshot = await session.get(ReportSnapshot, job_id)
        document_id = snapshot.document_id
        snapshot.expires_at = utcnow() - timedelta(days=1)
    async with _factory(engine)() as session:
        assert await reports_service.purge_expired_snapshots(session, tenant) == 0
        assert await session.get(ReportSnapshot, job_id) is not None
    app.dependency_overrides[get_current_principal] = lambda: replace(
        principal, allowed_warehouse_ids=(outside,)
    )
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 404
    app.dependency_overrides[get_current_principal] = lambda: replace(
        principal, allowed_warehouse_ids=(allowed,)
    )
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 200
    async with _factory(engine)() as session:
        await DocumentService(session, tenant).delete_document(
            document_id, reason="Retention regression cleanup"
        )
    async with _factory(engine)() as session:
        snapshot = await session.get(ReportSnapshot, job_id)
        assert snapshot.document_id is None
        assert await reports_service.purge_expired_snapshots(session, tenant) == 1
        assert await session.get(ReportSnapshot, job_id) is None


@pytest.mark.parametrize("legacy_attachment", [False, True])
async def test_orphan_valuation_document_stays_hidden_from_restricted_readers(
    client, engine, tenant, principal, legacy_attachment
):
    from src.main import app
    from src.modules.reports.models import ReportSnapshot

    (allowed, _), _ = await _stock_two_warehouses(engine, tenant)
    restricted = replace(principal, allowed_warehouse_ids=(allowed,))
    app.dependency_overrides[get_current_principal] = lambda: restricted
    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202
    job_id = uuid.UUID(accepted.json()["data"]["job_id"])
    await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=str(job_id),
    )
    async with _factory(engine)() as session, session.begin():
        snapshot = await session.get(ReportSnapshot, job_id)
        document = await session.get(Document, snapshot.document_id)
        document_id = document.id
        # Model the old cleanup behavior and artifacts created before attachments.
        if legacy_attachment:
            document.entity_type = None
            document.entity_id = None
        document.file_name = f"inventory-valuation-orphan-{job_id}.xlsx"
        file_name = document.file_name
        await session.delete(snapshot)
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 404
    listed = await client.get(
        f"{_prefix()}/documents", params={"search": file_name, "with_total": "true"}
    )
    assert listed.status_code == 200
    assert listed.json()["data"] == []
    assert listed.json()["pagination"]["total"] == 0
    for reader in (principal, replace(restricted, is_superuser=True)):
        app.dependency_overrides[get_current_principal] = lambda current=reader: current
        assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 200
        assert (
            await client.get(f"{_prefix()}/documents/{document_id}/download")
        ).status_code == 200


@pytest.mark.parametrize("scope", [None, [], "invalid"])
async def test_restricted_owner_cannot_read_untracked_or_full_company_valuation(
    client, engine, tenant, principal, scope
):
    from src.main import app
    from src.modules.reports.models import ReportSnapshot

    (allowed, _), _ = await _stock_two_warehouses(engine, tenant)
    exporter = replace(
        principal,
        allowed_warehouse_ids=(allowed,),
        permissions=frozenset(
            {str(Permission.INVENTORY_VALUATION_VIEW), str(Permission.REPORT_EXPORT)}
        ),
    )
    app.dependency_overrides[get_current_principal] = lambda: exporter
    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202
    job_id = uuid.UUID(accepted.json()["data"]["job_id"])
    await _run_worker(
        report_tasks.inventory_valuation,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
        snapshot_id=str(job_id),
    )
    async with _factory(engine)() as session, session.begin():
        snapshot = await session.get(ReportSnapshot, job_id)
        document_id = snapshot.document_id
        scope_key = reports_service.VALUATION_WAREHOUSE_SCOPE_PARAMETER
        parameters = dict(snapshot.parameters)
        if scope is None:
            parameters.pop(scope_key)
        else:
            parameters[scope_key] = scope
        snapshot.parameters = parameters
    assert (await client.get(reports_service.job_status_url(job_id))).status_code == 404
    assert (await client.get(reports_service.job_download_url(job_id))).status_code == 404
    app.dependency_overrides[get_current_principal] = lambda: replace(
        principal, allowed_warehouse_ids=(allowed,)
    )
    assert (await client.get(f"{_prefix()}/documents/{document_id}")).status_code == 404
    assert (await client.get(f"{_prefix()}/documents/{document_id}/download")).status_code == 404
    listed = await client.get(
        f"{_prefix()}/documents",
        params={"entity_type": report_tasks.VALUATION_ENTITY, "entity_id": str(job_id)},
    )
    assert listed.status_code == 200
    assert listed.json()["data"] == []
    assert listed.json()["pagination"]["total"] == 0


@pytest.mark.parametrize("wrong_field", ["report_key", "task"])
async def test_valuation_worker_refuses_a_job_without_finance_export_provenance(
    client, engine, tenant, wrong_field
):
    from src.core.clock import utcnow
    from src.core.errors import BusinessRuleError
    from src.modules.finance.service import VALUATION_TASK_NAME
    from src.modules.reports.models import ReportSnapshot

    parameters = {
        "task": VALUATION_TASK_NAME if wrong_field != "task" else "reports.generate_report",
        reports_service.VALUATION_WAREHOUSE_SCOPE_PARAMETER: [],
    }
    async with _factory(engine)() as session, session.begin():
        snapshot = reports_service.new_job_snapshot(
            ReportKey.STOCK_AVAILABILITY
            if wrong_field == "report_key"
            else ReportKey.INVENTORY_VALUATION,
            parameters,
            now=utcnow(),
        )
        ReportSnapshotRepository(session, tenant).add(snapshot)
        await session.flush()
        job_id = snapshot.id
        before = await session.scalar(select(func.count()).select_from(Document))
    with pytest.raises(BusinessRuleError, match="not a finance valuation export"):
        await _run_worker(
            report_tasks.inventory_valuation,
            company_id=str(tenant.company_id),
            user_id=str(tenant.user_id),
            snapshot_id=str(job_id),
        )
    async with _factory(engine)() as session:
        snapshot = await session.get(ReportSnapshot, job_id)
        assert snapshot.status == "FAILED"
        assert snapshot.document_id is None
        assert await session.scalar(select(func.count()).select_from(Document)) == before


async def test_legacy_pending_valuation_refuses_inactive_creator(client, engine, tenant):
    from src.core.errors import BusinessRuleError
    from src.modules.reports.models import ReportSnapshot
    from src.modules.users.models import User

    accepted = await client.post(f"{_prefix()}/finance/inventory-valuation/export")
    assert accepted.status_code == 202
    job_id = accepted.json()["data"]["job_id"]
    scope_key = reports_service.VALUATION_WAREHOUSE_SCOPE_PARAMETER
    async with _factory(engine)() as session, session.begin():
        snapshot = await session.get(ReportSnapshot, uuid.UUID(job_id))
        snapshot.parameters = {
            key: value for key, value in snapshot.parameters.items() if key != scope_key
        }
        creator = await session.get(User, tenant.user_id)
        creator.is_active = False
    with pytest.raises(BusinessRuleError, match="active internal creator"):
        await _run_worker(
            report_tasks.inventory_valuation,
            company_id=str(tenant.company_id),
            user_id=str(tenant.user_id),
            snapshot_id=job_id,
        )
    async with _factory(engine)() as session:
        snapshot = await session.get(ReportSnapshot, uuid.UUID(job_id))
        assert snapshot.status == "FAILED"
        assert snapshot.document_id is None
    async with _factory(engine)() as session, session.begin():
        creator = await session.get(User, tenant.user_id)
        creator.is_active = True


async def test_the_packing_list_is_a_job_that_files_a_pdf(
    client: AsyncClient, engine: Any, tenant: TenantContext
) -> None:
    order_id = await _packing_in_progress(engine, tenant)

    accepted = await client.get(f"{_prefix()}/packing/{order_id}/packing-list")
    assert accepted.status_code == 202, accepted.text
    handed_off = accepted.json()["data"]
    job_id = handed_off["job_id"]
    assert handed_off["status_url"] == reports_service.job_status_url(job_id)

    _assert_queued(
        await _job(client, handed_off["status_url"]), job_id=job_id, report_key="PACKING_LIST"
    )

    outcome = await _run_worker(
        report_tasks.generate_packing_list,
        packing_order_id=str(order_id),
        snapshot_id=job_id,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
    )
    assert outcome["status"] == "COMPLETED", outcome

    document_id = _assert_completed(
        await _job(client, handed_off["status_url"]), job_id=job_id, report_key="PACKING_LIST"
    )
    document = await _filed_document(engine, document_id)
    assert document.company_id == tenant.company_id
    assert document.category == DocumentCategory.REPORT
    assert document.entity_type == report_tasks.PACKING_LIST_ENTITY
    assert document.entity_id == order_id
    assert document.file_name.endswith(".pdf")
    assert resolve_storage_path(document.storage_key).is_file()


async def test_sales_invoice_pdf_job_files_permanent_document_and_dispatch_link(
    client, engine, tenant
):
    from src.core.clock import today
    from src.modules.allocation import service as allocation_service
    from src.modules.allocation.models import DemandType
    from src.modules.dispatch import service as dispatch_service
    from src.modules.dispatch.models import Dispatch
    from src.modules.dispatch.schemas import DispatchCreate, DispatchLineIn, DispatchRequest
    from src.modules.finance import sales_invoice_service
    from src.modules.finance.models import SalesInvoice
    from src.modules.finance.sales_invoice_schemas import SalesInvoiceCreate
    from src.modules.sales.models import SalesOrder, SalesOrderLine, SalesOrderStatus

    async with _factory(engine)() as session:
        uom = await seed_support.seed_uom(session, company_id=tenant.company_id)
        item = await seed_support.seed_item(
            session,
            company_id=tenant.company_id,
            code=_code("FG"),
            base_uom_id=uom.id,
            can_sell=True,
            hsn_code="8302",
        )
        warehouse = await seed_support.seed_warehouse(
            session, company_id=tenant.company_id, code=_code("WH")
        )
        customer = await seed_support.seed_customer(
            session, company_id=tenant.company_id, credit_days=30
        )
        order = SalesOrder(
            company_id=tenant.company_id,
            code=_code("SO"),
            customer_id=customer.id,
            warehouse_id=warehouse.id,
            order_date=today(),
            required_date=today(),
            status=SalesOrderStatus.ALLOCATED,
            is_interstate=True,
        )
        session.add(order)
        await session.flush()
        line = SalesOrderLine(
            company_id=tenant.company_id,
            sales_order_id=order.id,
            line_number=1,
            item_id=item.id,
            uom_id=uom.id,
            quantity=Decimal("10"),
            unit_price=Decimal("100"),
            tax_rate=Decimal("18"),
            allocated_quantity=Decimal("10"),
        )
        session.add(line)
        await session.flush()
        await inventory_service.post_movement(
            session,
            tenant,
            item_id=item.id,
            warehouse_id=warehouse.id,
            quantity=Decimal("10"),
            transaction_type=TransactionType.OPENING_BALANCE,
            reference_type="opening",
            reference_id=None,
            reference_code="OPEN",
            reason="PDF fixture stock",
            unit_cost=Decimal("50"),
        )
        reservation = await allocation_service.reserve(
            session,
            tenant,
            item_id=item.id,
            warehouse_id=warehouse.id,
            quantity=Decimal("10"),
            demand_type=DemandType.SALES_ORDER,
            demand_id=order.id,
            demand_line_id=line.id,
            reason="PDF fixture reservation",
        )
        line.reservation_id = reservation.id
        await session.flush()
        dispatch = await dispatch_service.create_dispatch(
            session,
            tenant,
            DispatchCreate(
                dispatch_type="SALES",
                dispatch_date=today(),
                customer_id=customer.id,
                warehouse_id=warehouse.id,
                sales_order_id=order.id,
                lines=[
                    DispatchLineIn(
                        item_id=item.id, sales_order_line_id=line.id, quantity=Decimal("10")
                    )
                ],
            ),
        )
        await dispatch_service.mark_ready(
            session, tenant, dispatch.id, reason="Ready for invoice PDF"
        )
        await dispatch_service.dispatch_consignment(
            session, tenant, dispatch.id, DispatchRequest(reason="Goods leave for PDF fixture")
        )
        invoice = await sales_invoice_service.issue_sales_invoice(
            session, tenant, SalesInvoiceCreate(dispatch_id=dispatch.id)
        )
        invoice_id, dispatch_id = invoice.id, dispatch.id
    accepted = await client.post(f"{_prefix()}/finance/sales-invoices/{invoice_id}/pdf")
    assert accepted.status_code == 202, accepted.text
    handed_off = accepted.json()["data"]
    job_id = handed_off["job_id"]
    _assert_queued(
        await _job(client, handed_off["status_url"]), job_id=job_id, report_key="SALES_INVOICE"
    )
    outcome = await _run_worker(
        report_tasks.generate_sales_invoice,
        sales_invoice_id=str(invoice_id),
        snapshot_id=job_id,
        company_id=str(tenant.company_id),
        user_id=str(tenant.user_id),
    )
    assert outcome["status"] == "COMPLETED", outcome
    document_id = _assert_completed(
        await _job(client, handed_off["status_url"]), job_id=job_id, report_key="SALES_INVOICE"
    )
    document = await _filed_document(engine, document_id)
    assert document.entity_type == "sales_invoice"
    assert document.entity_id == invoice_id
    assert document.expires_at is None
    assert resolve_storage_path(document.storage_key).read_bytes().startswith(b"%PDF")
    async with _factory(engine)() as reader:
        row = await reader.get(SalesInvoice, invoice_id)
        dispatched = await reader.get(Dispatch, dispatch_id)
        assert row.pdf_document_id == document.id
        assert dispatched.invoice_document_id == document.id
